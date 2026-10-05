import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_proposal_docx(output_path):
    doc = Document()

    # Set page margins to standard 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Color definitions
    NAVY = RGBColor(30, 41, 59)      # #1E293B
    INDIGO = RGBColor(79, 70, 229)   # #4F46E5
    SLATE = RGBColor(71, 85, 105)    # #475569
    DARK = RGBColor(15, 23, 42)      # #0F172A
    EMERALD = RGBColor(16, 185, 129) # #10B981

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = DARK
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_custom_heading(text, level=1):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        if level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = INDIGO
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif level == 2:
            run.font.size = Pt(13)
            run.font.color.rgb = NAVY
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif level == 3:
            run.font.size = Pt(11.5)
            run.font.color.rgb = INDIGO
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        return p

    # --- COVER / HEADER BANNER ---
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_cell = header_table.cell(0, 0)
    header_cell.width = Inches(6.5)
    set_cell_background(header_cell, "1E293B")
    set_cell_margins(header_cell, top=240, bottom=240, left=240, right=240)

    p_badge = header_cell.paragraphs[0]
    p_badge.paragraph_format.space_after = Pt(2)
    r_badge = p_badge.add_run("PROPOSAL PENAWARAN & SPESIFIKASI TEKNIS")
    r_badge.font.size = Pt(9.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(165, 180, 252) # Light Indigo

    p_title = header_cell.add_paragraph()
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("LEGAL MANAGEMENT SYSTEM (LMS) ENTERPRISE")
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = header_cell.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(0)
    r_sub = p_sub.add_run("Modul Tambahan Add-on ERP Terintegrasi Single Sign-On (SSO) & HRIS User Directory")
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(226, 232, 240)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- SUMMARY TABLE ---
    sum_table = doc.add_table(rows=6, cols=2)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sum_table.autofit = False

    table_data = [
        ("Nama Modul", "Enterprise Legal Management System (LMS) Add-on"),
        ("Tipe Pengadaan", "Modul Tambahan (Add-on Module) untuk Ekosistem ERP Eksisting"),
        ("Model Integrasi", "Single Sign-On (SSO OIDC / OAuth2 / SAML) + Sync User Directory HRIS"),
        ("Arsitektur Sistem", "Vue.js 3, Tailwind CSS, Lucide Icons, REST API Architecture"),
        ("Total Nilai Investasi", "Rp 40.000.000,- (Empat Puluh Juta Rupiah) — Nett / All-In"),
        ("Masa Garansi & SLA", "90 Hari Kalender (3 Bulan) Free Support, Bug Fixing & Pendampingan UAT")
    ]

    for i, (k, v) in enumerate(table_data):
        cell_k = sum_table.cell(i, 0)
        cell_v = sum_table.cell(i, 1)
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        set_cell_background(cell_k, "F8FAFC")
        set_cell_background(cell_v, "FFFFFF" if i % 2 == 0 else "F8FAFC")
        set_cell_margins(cell_k, top=70, bottom=70, left=100, right=100)
        set_cell_margins(cell_v, top=70, bottom=70, left=100, right=100)

        pk = cell_k.paragraphs[0]
        rk = pk.add_run(k)
        rk.font.bold = True
        rk.font.size = Pt(10)
        rk.font.color.rgb = NAVY

        pv = cell_v.paragraphs[0]
        rv = pv.add_run(v)
        rv.font.size = Pt(10)
        if "Rp 40.000.000" in v:
            rv.font.bold = True
            rv.font.color.rgb = INDIGO

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 1: LATAR BELAKANG ---
    add_custom_heading("1. LATAR BELAKANG & TUJUAN INTEGRASI", 1)
    p = doc.add_paragraph(
        "Perusahaan telah memiliki infrastruktur ERP yang matang untuk tata kelola pengguna (User Management) "
        "dan Autentikasi Tunggal (Single Sign-On / SSO). Untuk melengkapi proses bisnis korporasi dalam aspek hukum, "
        "perlindungan risiko kontrak, kepatuhan perizinan, dan kualifikasi lelang tender, diimplementasikan modul "
        "Legal Management System (LMS) sebagai subsistem terintegrasi."
    )
    
    p = doc.add_paragraph()
    r = p.add_run("Keuntungan Utama Integrasi Modul Legal ke ERP Eksisting:")
    r.font.bold = True

    benefits = [
        ("Single Identity (SSO)", "Pengguna tidak perlu mengingat kredensial baru. Akses modul legal langsung terafiliasi dengan akun login ERP."),
        ("Role-Based Access Control (RBAC)", "Pemetaan hak akses (General Counsel, Legal Specialist, Contract Drafter, Auditor, Direksi) otomatis disinkronkan dari jabatan dan unit kerja di ERP HRIS."),
        ("Audit Trail Terpusat", "Seluruh tindakan penambahan, pengubahan, approval, dan penghapusan dokumen legal tercatat dengan identitas resmi karyawan ERP."),
        ("End-to-End Workflow", "Integrasi konversi proses bisnis, mulai dari tender yang dimenangkan langsung dikonversi menjadi Kontrak Korporasi pada modul legal dan ERP Keuangan.")
    ]
    for title, desc in benefits:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        rt = bp.add_run(f"{title}: ")
        rt.bold = True
        rt.font.color.rgb = NAVY
        bp.add_run(desc)

    # --- SECTION 2: DETAIL FITUR ---
    add_custom_heading("2. DETAIL LENGKAP FITUR MODUL LEGAL (14 SUB-MODUL)", 1)
    doc.add_paragraph(
        "Modul Legal Management System Enterprise mencakup 14 sub-modul operasional yang siap digunakan dan telah dilengkapi dengan kapabilitas CRUD lengkap (Create, Read, Update, Delete) serta penyesuaian regulasi Indonesia:"
    )

    features = [
        ("Fitur 1: Executive Dashboard & Legal Intelligence Analytics",
         "Pusat pemantauan visual metrik kesehatan legal perusahaan secara real-time. Menampilkan ringkasan kontrak aktif, nilai komitmen finansial, peringatan kontrak kedaluwarsa (<60 hari), Tender Readiness Score, eksposur sengketa perdata/BANI, dan kalender sidang."),
        
        ("Fitur 2: Manajemen Permohonan Layanan Hukum (Legal Service Request & SLA)",
         "Sistem tiket internal bagi seluruh unit kerja di ERP untuk mengajukan kebutuhan kajian kontrak, somasi, opini hukum, atau verifikasi legalitas mitra bisnis dengan penomoran otomatis (REQ-2026-XXXX) dan pelacakan SLA."),
        
        ("Fitur 3: Contract Lifecycle Management (CLM) & Manajemen Adendum",
         "Tata kelola siklus hidup kontrak bisnis mulai dari draf, negosiasi klausul, formulir Edit Kontrak interaktif, manajemen riwayat adendum perpanjangan/eskalasi nilai, proteksi penghapusan kontrak, serta ekspor rekapitulasi data register ke format CSV/Excel."),
        
        ("Fitur 4: Bank Dokumen Kualifikasi Pengadaan Lelang (Tender Management Module)",
         "Modul percepatan keikutsertaan lelang pengadaan (Konstruksi, Tambang, Energi, Migas, Umum). Dilengkapi Gap Analysis Dokumen Persyaratan, Evaluasi Go/No-Go 6 dimensi, Analisis Margin & HPS (<80% HPS alert), 4-Gate Review, Kunci Submission SHA-256, dan konversi lelang menang (WON) menjadi Kontrak Korporasi."),
        
        ("Fitur 5: Manajemen Perizinan Berusaha (OSS RBA, PB-UMKU & IUP)",
         "Pengelolaan seluruh izin operasional, lingkungan, ketenagalistrikan, dan pertambangan perseroan. Dilengkapi formulir Edit Izin, pemantauan masa berlaku, instansi penerbit (ESDM, BKPM, KLHK, PUPR), dan indikator visual jatuh tempo."),
        
        ("Fitur 6: Matriks Kepatuhan Hukum & Audit Regulasi (Compliance Tracker)",
         "Pengawasan kepatuhan mandatori berkala terhadap regulasi RI (Pelaporan LKPM BKPM triwulanan/semesteran, RKL-RPL KLHK, WLKP Ketenagakerjaan, RKAB Minerba) dengan fitur Tambah/Edit/Hapus kewajiban dan tingkat risiko."),
        
        ("Fitur 7: Manajemen Sengketa, Litigasi & Arbitrase (Disputes Tracker)",
         "Sentralisasi penanganan perkara sengketa hukum vs pihak lawan, lawyer penanggung jawab, pengadilan (PN/PT/MA/BANI/SIAC), nilai klaim eksposur kerugian, riwayat persidangan, dan update status berkekuatan hukum tetap (Inkracht)."),
        
        ("Fitur 8: Legal Due Diligence (LDD) & M&A Audit Checklist",
         "Toolkit audit kepatuhan hukum untuk transaksi M&A, akuisisi saham, joint venture, atau audit internal unit bisnis. Dilengkapi checklist multi-aspek dengan interaktif toggle status audit (OK / FLAG / N/A), tambah butir audit kustom, dan hapus proyek LDD."),
        
        ("Fitur 9: Penyusunan & Alur Persetujuan Legal Opinion (Kajian Hukum)",
         "Pembuatan nota kajian hukum terstruktur (Latar Belakang, Dasar Hukum & Analisis Yuridis, Kesimpulan, Rekomendasi Mitigasi Risiko) beserta alur persetujuan General Counsel / Head of Legal."),
        
        ("Fitur 10: Perpustakaan Klausul Standar & Interactive Template Generator",
         "Koleksi klausul baku teruji (Arbitrase BANI, FIDIC, batasan tanggung jawab, force majeure, kerahasiaan NDA) dengan fitur Tambah/Edit/Hapus klausul, tombol salin instan, dan Generator Naskah interaktif berbasis merge fields."),
        
        ("Fitur 11: Digital Legal & Tender Vault (Arsip Berkas Terenkripsi)",
         "Penyimpanan digital terpusat berklasifikasi ganda (Arsip Akta Notaris, SK Kemenkumham, Sertifikat Tanah HGB vs Bank Dokumen Kualifikasi NIB/SBU/ISO/KAP) dengan kapabilitas Upload, Edit Metadata, Unduh Salinan, dan Tautkan ke Tender."),
        
        ("Fitur 12: Buku Register Surat & Korespondensi Legal (In/Out Letters)",
         "Pencatatan digital seluruh korespondensi resmi, somasi, surat masuk kementerian, dan surat peringatan dengan penomoran standar korporasi otomatis, tracking deadline, dan lampiran berkas digital."),
        
        ("Fitur 13: Database Regulasi Positif & Yurisprudensi (JDIH Knowledge Base)",
         "Kompilasi regulasi perundang-undangan sektoral (UU, PP, Permen ESDM, Permen BKPM, KLHK) dan Putusan MA dengan ringkasan ketentuan pokok, pasal kritis, tautan JDIH resmi, dan fitur Tambah/Edit/Hapus regulasi."),
        
        ("Fitur 14: Integrasi Role-Based Access Control (RBAC) & Audit Trail ERP",
         "Matriks 13 hak akses granular terintegrasi dengan akun ERP, dukungan custom role, serta log audit trail yang merekam seluruh aksi (Create, Update, Delete, Approve, Export) secara persistent.")
    ]

    for fname, fdesc in features:
        add_custom_heading(fname, 3)
        p = doc.add_paragraph(fdesc)
        p.paragraph_format.space_after = Pt(4)

    # --- SECTION 3: ARSITEKTUR INTEGRASI ---
    add_custom_heading("3. ARSITEKTUR INTEGRASI SINGLE SIGN-ON (SSO) & ERP", 1)
    doc.add_paragraph(
        "Modul Legal Management System dirancang dengan standar arsitektur RESTful & decoupled frontend yang sangat mudah dihubungkan dengan SSO Gateway ERP:"
    )

    sso_steps = [
        ("Protokol Autentikasi Standar", "Mendukung OAuth 2.0 / OpenID Connect (OIDC), SAML 2.0, serta JWT Bearer Token yang diterbitkan oleh SSO ERP."),
        ("User Directory Synchronization", "Data identitas (Nama, NIK, Email, Departemen, Jabatan) diambil otomatis saat sesi pertama login dan diperbarui secara berkala via REST API."),
        ("Role & Capability Mapping", "Jabatan di ERP (misal: 'Legal Specialist', 'Finance Manager', 'Operations Director') secara otomatis dipetakan ke role matriks wewenang LMS."),
        ("Multi-Entity Context", "Mendukung pengelolaan multi-perusahaan (Holding dan Anak Perusahaan) sesuai entitas akun pengguna di ERP.")
    ]
    for stitle, sdesc in sso_steps:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        rt = bp.add_run(f"{stitle}: ")
        rt.bold = True
        rt.font.color.rgb = NAVY
        bp.add_run(sdesc)

    # --- SECTION 4: RINCIAN PENAWARAN BIAYA ---
    add_custom_heading("4. PENAWARAN BIAYA & RINCIAN INVESTASI", 1)
    p_price_banner = doc.add_paragraph()
    r_pb = p_price_banner.add_run("TOTAL NILAI INVESTASI: Rp 40.000.000,- (EMPAT PULUH JUTA RUPIAH) NETTO")
    r_pb.bold = True
    r_pb.font.size = Pt(12)
    r_pb.font.color.rgb = INDIGO
    p_price_banner.paragraph_format.space_after = Pt(6)

    price_table = doc.add_table(rows=8, cols=3)
    price_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    price_table.autofit = False

    headers = ["No", "Komponen Pekerjaan & Deliverables", "Nilai Investasi (IDR)"]
    for col_idx, htext in enumerate(headers):
        cell = price_table.cell(0, col_idx)
        cell.width = Inches(0.6) if col_idx == 0 else (Inches(4.4) if col_idx == 1 else Inches(1.5))
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(htext)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    price_rows = [
        ("1", "Lisensi Source Code & Modul Penuh LMS Enterprise\n• Seluruh 14 Sub-Modul Fitur Legal lengkap & responsive UI\n• Database Storage Schema, State Management & Validasi Form", "Rp 18.000.000,-"),
        ("2", "Konfigurasi & Integrasi Single Sign-On (SSO) ERP\n• Integrasi protokol OAuth2/OIDC/SAML ke Identity Provider ERP\n• Sinkronisasi Data Karyawan (HRIS Directory) & Role Mapping Otomatis", "Rp 8.000.000,-"),
        ("3", "Setup Environment, Deployment & Database Seeding\n• Instalasi & integrasi ke server aplikasi ERP\n• Seeding data awal (Klausul BANI/FIDIC, Master Sektor, Regulasi JDIH, Matriks Izin)", "Rp 6.000.000,-"),
        ("4", "Testing, Quality Assurance & UAT (User Acceptance Testing)\n• Uji fungsional seluruh alur CRUD & Workflow Approval\n• Uji keamanan autentikasi SSO & proteksi data vault", "Rp 4.000.000,-"),
        ("5", "Dokumentasi Teknis, Panduan Pengguna & Pelatihan (Training)\n• Manual Book Admin & User Guide (PDF)\n• Sesi pelatihan daring / luring untuk Tim Legal & Admin ERP", "Rp 2.000.000,-"),
        ("6", "Garansi, Pemeliharaan & SLA Support (3 Bulan)\n• Pendampingan pasca Go-Live, Bug Fixing & Konsultasi Teknis", "Rp 2.000.000,-"),
        ("", "TOTAL INVESTASI KESELURUHAN (ALL-IN)", "Rp 40.000.000,-")
    ]

    for row_idx, (c0, c1, c2) in enumerate(price_rows, start=1):
        cell0 = price_table.cell(row_idx, 0)
        cell1 = price_table.cell(row_idx, 1)
        cell2 = price_table.cell(row_idx, 2)
        cell0.width = Inches(0.6)
        cell1.width = Inches(4.4)
        cell2.width = Inches(1.5)

        is_total = (row_idx == 7)
        bg_color = "EEF2FF" if is_total else ("F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        set_cell_background(cell0, bg_color)
        set_cell_background(cell1, bg_color)
        set_cell_background(cell2, bg_color)
        set_cell_margins(cell0, top=70, bottom=70, left=80, right=80)
        set_cell_margins(cell1, top=70, bottom=70, left=100, right=100)
        set_cell_margins(cell2, top=70, bottom=70, left=80, right=80)

        p0 = cell0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(c0)
        r0.font.bold = is_total

        p1 = cell1.paragraphs[0]
        r1 = p1.add_run(c1)
        r1.font.bold = is_total
        if is_total:
            r1.font.color.rgb = INDIGO

        p2 = cell2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(c2)
        r2.font.bold = True
        if is_total:
            r2.font.color.rgb = INDIGO
            r2.font.size = Pt(11)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 5: TIMELINE ---
    add_custom_heading("5. JADWAL PELAKSANAAN & TIMELINE IMPLEMENTASI", 1)
    doc.add_paragraph("Total durasi pengerjaan dan integrasi diestimasikan selama 4 (Empat) Minggu:")

    timeline_table = doc.add_table(rows=5, cols=3)
    timeline_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    timeline_table.autofit = False

    t_headers = ["Minggu", "Tahapan Kerja", "Target Output"]
    for col_idx, htext in enumerate(t_headers):
        cell = timeline_table.cell(0, col_idx)
        cell.width = Inches(1.0) if col_idx == 0 else (Inches(3.2) if col_idx == 1 else Inches(2.3))
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(htext)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    t_rows = [
        ("Minggu 1", "Analisis spesifikasi SSO ERP eksisting & Penyesuaian skema autentikasi dan Role Mapping", "Dokumen Konfigurasi SSO disepakati"),
        ("Minggu 2", "Integrasi modul frontend & API backend SSO serta Setup environment staging / deployment", "Modul terpasang di server staging"),
        ("Minggu 3", "Database initial seeding & Pelaksanaan User Acceptance Testing (UAT) bersama Tim Legal", "Berita Acara UAT & Penyesuaian Feedback"),
        ("Minggu 4", "Pelatihan Pengguna (User Training Session), Go-Live Production & Serah Terima (BAST)", "Modul aktif di Production & BAST diterbitkan")
    ]

    for row_idx, (c0, c1, c2) in enumerate(t_rows, start=1):
        cell0 = timeline_table.cell(row_idx, 0)
        cell1 = timeline_table.cell(row_idx, 1)
        cell2 = timeline_table.cell(row_idx, 2)
        cell0.width = Inches(1.0)
        cell1.width = Inches(3.2)
        cell2.width = Inches(2.3)

        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        set_cell_background(cell0, bg)
        set_cell_background(cell1, bg)
        set_cell_background(cell2, bg)
        set_cell_margins(cell0, top=70, bottom=70, left=80, right=80)
        set_cell_margins(cell1, top=70, bottom=70, left=100, right=100)
        set_cell_margins(cell2, top=70, bottom=70, left=100, right=100)

        p0 = cell0.paragraphs[0]
        r0 = p0.add_run(c0)
        r0.font.bold = True

        p1 = cell1.paragraphs[0]
        p1.add_run(c1)

        p2 = cell2.paragraphs[0]
        r2 = p2.add_run(c2)
        r2.font.bold = True
        r2.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 6: KETENTUAN TERMIN PEMBAYARAN ---
    add_custom_heading("6. KETENTUAN TERMIN PEMBAYARAN", 1)
    doc.add_paragraph("Pembayaran dilakukan secara bertahap melalui transfer bank resmi:")

    term_table = doc.add_table(rows=4, cols=4)
    term_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    term_table.autofit = False

    term_headers = ["Termin", "Porsi", "Nilai (IDR)", "Kondisi Penagihan"]
    for col_idx, htext in enumerate(term_headers):
        cell = term_table.cell(0, col_idx)
        cell.width = Inches(1.2) if col_idx == 0 else (Inches(0.8) if col_idx == 1 else (Inches(1.5) if col_idx == 2 else Inches(3.0)))
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(htext)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    term_rows = [
        ("Termin I (DP)", "40%", "Rp 16.000.000,-", "Saat penandatanganan Kontrak Kerja Sama / SPK"),
        ("Termin II (UAT)", "40%", "Rp 16.000.000,-", "Setelah integrasi SSO selesai & lulus uji coba UAT"),
        ("Termin III (Final)", "20%", "Rp 8.000.000,-", "Setelah Go-Live, Training selesai & TTD BAST")
    ]

    for row_idx, (c0, c1, c2, c3) in enumerate(term_rows, start=1):
        cell0 = term_table.cell(row_idx, 0)
        cell1 = term_table.cell(row_idx, 1)
        cell2 = term_table.cell(row_idx, 2)
        cell3 = term_table.cell(row_idx, 3)

        cell0.width = Inches(1.2)
        cell1.width = Inches(0.8)
        cell2.width = Inches(1.5)
        cell3.width = Inches(3.0)

        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        set_cell_background(cell0, bg)
        set_cell_background(cell1, bg)
        set_cell_background(cell2, bg)
        set_cell_background(cell3, bg)
        set_cell_margins(cell0, top=70, bottom=70, left=80, right=80)
        set_cell_margins(cell1, top=70, bottom=70, left=80, right=80)
        set_cell_margins(cell2, top=70, bottom=70, left=80, right=80)
        set_cell_margins(cell3, top=70, bottom=70, left=80, right=80)

        p0 = cell0.paragraphs[0]
        r0 = p0.add_run(c0)
        r0.font.bold = True

        p1 = cell1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run(c1)

        p2 = cell2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(c2)
        r2.font.bold = True

        p3 = cell3.paragraphs[0]
        p3.add_run(c3)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # --- SECTION 7: LEMBAR PENGESAHAN ---
    add_custom_heading("7. LEMBAR PERSETUJUAN & KONFIRMASI", 1)
    doc.add_paragraph("Dokumen penawaran dan spesifikasi modul ini disetujui bersama oleh para pihak:")

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    sig_left = sig_table.cell(0, 0)
    sig_right = sig_table.cell(0, 1)
    sig_left.width = Inches(3.25)
    sig_right.width = Inches(3.25)
    set_cell_background(sig_left, "F8FAFC")
    set_cell_background(sig_right, "F8FAFC")
    set_cell_margins(sig_left, top=140, bottom=140, left=140, right=140)
    set_cell_margins(sig_right, top=140, bottom=140, left=140, right=140)

    pl = sig_left.paragraphs[0]
    pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rl1 = pl.add_run("Disiapkan Oleh,\nTim Pengembang / Solution Architect\n\n\n\n\n\n")
    rl1.font.size = Pt(10)
    rl2 = pl.add_run("___________________________\n( Solution Architect Lead )")
    rl2.font.bold = True

    pr = sig_right.paragraphs[0]
    pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr1 = pr.add_run("Disetujui Oleh,\nDireksi / Tim Manajemen ERP\n\n\n\n\n\n")
    rr1.font.size = Pt(10)
    rr2 = pr.add_run("___________________________\n( Direktur / IT Project Lead )")
    rr2.font.bold = True

    # Save document
    doc.save(output_path)
    print(f"Successfully generated DOCX at {output_path}")

if __name__ == '__main__':
    create_proposal_docx("/Users/mac/Documents/antigravity/web-legal/DOKUMEN_PENAWARAN_DAN_SPESIFIKASI_MODUL_LEGAL.docx")
